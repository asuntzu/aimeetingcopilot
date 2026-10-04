#!/usr/bin/env python3
"""Print the plain text of a document (PDF, Word .docx, RTF, HTML, text) for the copilot's index."""
import re
import subprocess
import sys


def tidy(page):
    """Some PDFs extract one word per line; rejoin those into normal sentences."""
    lines = [l.strip() for l in page.splitlines()]
    words = [l for l in lines if l]
    if words and sum(len(l) for l in words) / len(words) < 15:
        return re.sub(r"\s+", " ", " ".join(words))
    return page


def extract(path):
    low = path.lower()
    if low.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n\n".join(tidy(p.extract_text() or "") for p in PdfReader(path).pages)
    if low.endswith(".docx"):
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if low.endswith(".rtf"):
        from striprtf.striprtf import rtf_to_text
        with open(path, encoding="utf-8", errors="ignore") as f:
            return rtf_to_text(f.read())
    if low.endswith((".doc", ".pages", ".odt")) and sys.platform == "darwin":
        return subprocess.run(["/usr/bin/textutil", "-convert", "txt", "-stdout", path],
                              capture_output=True, text=True).stdout
    with open(path, encoding="utf-8", errors="ignore") as f:
        text = f.read()
    if low.endswith((".html", ".htm")):
        text = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", text)
        text = re.sub(r"<[^>]+>", " ", text)
    return text


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stdout.write(extract(sys.argv[1]))
    except Exception as e:  # noqa: BLE001
        sys.stderr.write(f"extract: {e}\n")
        sys.exit(1)
